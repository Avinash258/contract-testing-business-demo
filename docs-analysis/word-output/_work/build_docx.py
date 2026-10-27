#!/usr/bin/env python3
"""Build Contract-Testing-Business-Demo-Complete-Guide.docx from docs-analysis MD files."""
from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
DOCS = ROOT.parent.parent
OUT = ROOT.parent / "Contract-Testing-Business-Demo-Complete-Guide.docx"
DIAG = ROOT / "diagrams"
META = json.loads((ROOT / "figures.json").read_text(encoding="utf-8"))

FILES = [
    ("00-README-INDEX.md", "Index and Discovery"),
    ("01-beginner-concepts.md", "Beginner Concepts"),
    ("02-architecture.md", "Architecture"),
    ("03-file-by-file-guide.md", "File-by-File Guide"),
    ("04-end-to-end-flow.md", "End-to-End Flow"),
    ("05-how-it-runs-explained.md", "How It Runs (Explained)"),
    ("06-glossary-and-qna.md", "Glossary and Q&A"),
]

MERMAID_RE = re.compile(r"```mermaid\s*\n(.*?)```", re.DOTALL)
CODE_RE = re.compile(r"```(\w*)\s*\n(.*?)```", re.DOTALL)
TABLE_SEP_RE = re.compile(r"^\|[\s:\-|]+\|$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
HR_RE = re.compile(r"^---+\s*$")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
ITALIC_RE = re.compile(r"(?<!\*)\*([^*]+)\*(?!\*)")
CODE_INLINE_RE = re.compile(r"`([^`]+)`")

# Map cross-file md refs to section titles
FILE_REF_MAP = {
    "00-README-INDEX.md": "Index and Discovery",
    "01-beginner-concepts.md": "Beginner Concepts",
    "02-architecture.md": "Architecture",
    "03-file-by-file-guide.md": "File-by-File Guide",
    "04-end-to-end-flow.md": "End-to-End Flow",
    "05-how-it-runs-explained.md": "How It Runs (Explained)",
    "06-glossary-and-qna.md": "Glossary and Q&A",
    "01-…": "Beginner Concepts",
    "02-…": "Architecture",
    "03-…": "File-by-File Guide",
    "04-…": "End-to-End Flow",
    "05-…": "How It Runs (Explained)",
    "06-…": "Glossary and Q&A",
}


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=None, mono=False):
    run.font.name = "Consolas" if mono else name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), run.font.name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    # remove existing shd
    for child in list(tcPr):
        if child.tag == qn("w:shd"):
            tcPr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "666666")
        borders.append(el)
    tcPr.append(borders)


def set_repeat_header(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instr)
    run._r.append(fldChar2)


def add_toc_field(paragraph):
    """Insert Word TOC field (updates when opened / Update Fields)."""
    run = paragraph.add_run()
    r = run._r

    def fld(char_type, text=None):
        if char_type:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), char_type)
            r.append(fc)
        if text is not None:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r.append(it)

    fld("begin")
    fld(None, ' TOC \\o "1-3" \\h \\z \\u ')
    fld("separate")
    # placeholder shown until Word refreshes
    t = OxmlElement("w:t")
    t.text = "Right-click → Update Field to refresh Table of Contents."
    r.append(t)
    fld("end")


def set_update_fields_on_open(doc):
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def add_horizontal_line(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)


def rewrite_file_refs(text: str) -> str:
    out = text
    for k, v in FILE_REF_MAP.items():
        # backtick paths
        out = out.replace(f"`{k}`", f"section “{v}”")
        out = out.replace(k, v)
    # docs-analysis/01-... style
    out = re.sub(
        r"docs-analysis/(0[0-6]-[a-z0-9\-]+\.md)",
        lambda m: f'section “{FILE_REF_MAP.get(m.group(1), m.group(1))}”',
        out,
    )
    return out


def add_formatted_runs(paragraph, text: str, base_size=11):
    text = rewrite_file_refs(text)
    # Convert markdown links to "text (target)" or section name
    def link_sub(m):
        label, href = m.group(1), m.group(2)
        if href.endswith(".md") or href in FILE_REF_MAP:
            name = FILE_REF_MAP.get(Path(href).name, FILE_REF_MAP.get(href, label))
            return f"{label} → {name}"
        return f"{label} ({href})"

    text = LINK_RE.sub(link_sub, text)

    # Tokenize bold / italic / inline code
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=base_size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=base_size - 1, mono=True, color=RGBColor(0x33, 0x33, 0x33))
            # light grey highlight
            rPr = run._r.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F0F0F0")
            rPr.append(shd)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=base_size, italic=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=base_size)


def style_table(table, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    # Preferred width ~ page content
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)
    # Table width 100%
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "pct")
    tblW.set(qn("w:w"), "5000")  # 100%
    # remove existing tblW
    for child in list(tblPr):
        if child.tag == qn("w:tblW"):
            tblPr.remove(child)
    tblPr.append(tblW)

    for i, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_borders(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    if run.font.size is None:
                        run.font.size = Pt(9)
                    run.font.name = "Calibri"
            if header and i == 0:
                shade_cell(cell, "1F4E79")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.bold = True
                        run.font.size = Pt(9)
        if header and i == 0:
            set_repeat_header(row)


def parse_table_lines(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        if TABLE_SEP_RE.match(line.strip()):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def add_code_block(doc, code: str, lang: str = ""):
    # One paragraph per line for monospace block with shading via paragraph border/shading
    for i, line in enumerate(code.splitlines() or [""]):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.3)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "F2F2F2")
        pPr.append(shd)
        run = p.add_run(line if line else " ")
        set_run_font(run, size=8.5, mono=True, color=RGBColor(0x22, 0x22, 0x22))
    # spacer
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_figure(doc, fig_meta: dict):
    png = DIAG / fig_meta["png"]
    if not png.exists():
        p = doc.add_paragraph(f"[Missing figure: {fig_meta['png']}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    # Fit to ~6.2 inches content width
    run.add_picture(str(png), width=Inches(6.2))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run(f"Figure {fig_meta['fig']}: {fig_meta['title']}")
    set_run_font(r, size=10, italic=True, color=RGBColor(0x33, 0x33, 0x33))


def setup_styles(doc: Document):
    styles = doc.styles
    # Normal
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for level, size in ((1, 18), (2, 14), (3, 12)):
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        style.paragraph_format.space_after = Pt(6)


def setup_header_footer(section, title: str):
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.clear()
    hr = hp.add_run(title)
    set_run_font(hr, size=9, italic=True, color=RGBColor(0x66, 0x66, 0x66))
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Page ")
    set_run_font(fr, size=9)
    add_page_number(fp)
    fr2 = fp.add_run("  |  Contract Testing Business Demo — Complete Guide")
    set_run_font(fr2, size=9, color=RGBColor(0x66, 0x66, 0x66))


def cover_page(doc: Document):
    for _ in range(3):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Contract Testing Business Demo")
    set_run_font(r, size=28, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("Complete Architecture and Flow Guide")
    set_run_font(r, size=18, italic=True, color=RGBColor(0x2E, 0x75, 0xB6))

    add_horizontal_line(doc)

    meta_lines = [
        ("Repository", "Avinash258/contract-testing-business-demo"),
        ("Prepared for", "QA / Test Architecture learners"),
        ("Document date", date.today().strftime("%d %B %Y")),
        ("Source", "docs-analysis/00 … 06 Markdown pack (explain-only analysis)"),
        ("Diagrams", f"{len(META['figures'])} Mermaid figures rendered to PNG"),
    ]
    for label, value in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run(f"{label}: ")
        set_run_font(r1, size=12, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=12)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run(
        "This document merges the docs-analysis Markdown files without rewriting their content. "
        "Mermaid diagrams are embedded as figures. Commands described in the source were not executed for the analysis."
    )
    set_run_font(r, size=10, italic=True, color=RGBColor(0x55, 0x55, 0x55))

    # page break after cover
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def toc_and_lod(doc: Document):
    h = doc.add_heading("Table of Contents", level=1)
    p = doc.add_paragraph()
    add_toc_field(p)
    # Static outline (always visible)
    doc.add_heading("Document outline (static)", level=2)
    for i, (fname, title) in enumerate(FILES, start=1):
        p = doc.add_paragraph(style="List Number")
        add_formatted_runs(p, f"**{title}** — source `{fname}`")

    doc.add_heading("List of Diagrams", level=1)
    for fig in META["figures"]:
        p = doc.add_paragraph(style="List Number")
        add_formatted_runs(
            p,
            f"**Figure {fig['fig']}:** {fig['title']} _(source: `{fig['file']}`, type: {fig['type']})_",
        )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def process_md_file(doc: Document, fname: str, section_title: str, fig_cursor: list):
    """fig_cursor is a one-element list holding next figure index (0-based into META['figures'])."""
    # Major section on new page (except we already broke before first)
    doc.add_heading(section_title, level=1)
    src_note = doc.add_paragraph()
    r = src_note.add_run(f"Source file: {fname}")
    set_run_font(r, size=9, italic=True, mono=True, color=RGBColor(0x66, 0x66, 0x66))

    text = (DOCS / fname).read_text(encoding="utf-8")
    # Split while keeping mermaid/code fences
    # Walk line by line for tables/lists/headings; detect fenced blocks
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]

        # Mermaid or code fence
        if line.strip().startswith("```"):
            lang = line.strip()[3:].strip()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1  # closing fence
            body = "\n".join(block)
            if lang.lower() == "mermaid":
                if fig_cursor[0] < len(META["figures"]):
                    add_figure(doc, META["figures"][fig_cursor[0]])
                    fig_cursor[0] += 1
                else:
                    doc.add_paragraph("[Figure index overflow]")
            else:
                add_code_block(doc, body, lang)
            continue

        # Table start
        if line.strip().startswith("|") and i + 1 < len(lines) and TABLE_SEP_RE.match(lines[i + 1].strip()):
            tlines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tlines.append(lines[i])
                i += 1
            rows = parse_table_lines(tlines)
            if not rows:
                continue
            cols = max(len(r) for r in rows)
            # normalize col counts
            for r in rows:
                while len(r) < cols:
                    r.append("")
            # Wide table heuristic
            wide = cols >= 6 or any(sum(len(c) for c in r) > 140 for r in rows)
            font_size = 8 if wide else 9

            table = doc.add_table(rows=len(rows), cols=cols)
            for ri, row in enumerate(rows):
                for ci, cell_text in enumerate(row):
                    cell = table.cell(ri, ci)
                    cell.text = ""
                    p = cell.paragraphs[0]
                    # strip md bold for simplicity then format
                    add_formatted_runs(p, cell_text, base_size=font_size)
            style_table(table, header=True)
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue

        # Heading
        hm = HEADING_RE.match(line)
        if hm:
            level = len(hm.group(1))
            title = hm.group(2).strip()
            # Skip duplicating top H1 that restates file title — keep content
            # Map H1 in file to H2 under our section H1; H2->H3; H3->H3 max for TOC depth
            mapped = min(level + 1, 3)
            # If original was H1 like "# 01 — ...", use as H2
            h = doc.add_heading(level=mapped)
            add_formatted_runs(h, title, base_size=18 if mapped == 1 else (14 if mapped == 2 else 12))
            for run in h.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            i += 1
            continue

        # HR
        if HR_RE.match(line.strip()):
            add_horizontal_line(doc)
            i += 1
            continue

        # Empty
        if not line.strip():
            i += 1
            continue

        # Blockquote
        if line.strip().startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            add_formatted_runs(p, line.strip().lstrip(">").strip())
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        # List
        if re.match(r"^[-*]\s+", line.strip()) or re.match(r"^\d+\.\s+", line.strip()):
            style = "List Number" if re.match(r"^\d+\.", line.strip()) else "List Bullet"
            content = re.sub(r"^([-*]|\d+\.)\s+", "", line.strip())
            p = doc.add_paragraph(style=style)
            add_formatted_runs(p, content)
            i += 1
            continue

        # Normal paragraph — gather continuation lines?
        p = doc.add_paragraph()
        add_formatted_runs(p, line.strip())
        i += 1

    # page break after section
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def main():
    doc = Document()
    # Page setup A4
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    setup_styles(doc)
    set_update_fields_on_open(doc)

    cover_page(doc)
    setup_header_footer(doc.sections[0], "Contract Testing Business Demo — Complete Guide")
    toc_and_lod(doc)

    fig_cursor = [0]
    for fname, title in FILES:
        num = fname[:2]
        process_md_file(doc, fname, f"{num} — {title}", fig_cursor)

    # Remove trailing empty page break if last para is only break — leave as is

    # Stats paragraph at end
    doc.add_heading("Build metadata", level=1)
    stats = doc.add_paragraph()
    add_formatted_runs(
        stats,
        f"Figures embedded: **{fig_cursor[0]}** / **{len(META['figures'])}**. "
        f"Fixes applied during diagram render: {len(META.get('fixes', []))} — "
        + (", ".join(META.get("fixes", [])) or "none")
        + ".",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")
    print(f"Size bytes: {OUT.stat().st_size}")
    print(f"Figures used: {fig_cursor[0]}")


if __name__ == "__main__":
    main()
