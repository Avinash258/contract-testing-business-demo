#!/usr/bin/env python3
"""Verify the generated DOCX quality."""
from __future__ import annotations

import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
DOCX = ROOT.parent / "Contract-Testing-Business-Demo-Complete-Guide.docx"
META = json.loads((ROOT / "figures.json").read_text(encoding="utf-8"))

EXPECTED_SECTIONS = [
    "00 — Index and Discovery",
    "01 — Beginner Concepts",
    "02 — Architecture",
    "03 — File-by-File Guide",
    "04 — End-to-End Flow",
    "05 — How It Runs (Explained)",
    "06 — Glossary and Q&A",
]


def main():
    doc = Document(str(DOCX))
    texts = []
    headings = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        texts.append(t)
        if p.style and p.style.name.startswith("Heading"):
            headings.append((p.style.name, t))

    full = "\n".join(texts)

    # Sections present
    section_status = []
    for s in EXPECTED_SECTIONS:
        section_status.append((s, s in full))

    # Images via relationships
    image_count = 0
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            image_count += 1

    # Also count in package
    with zipfile.ZipFile(DOCX) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        xml = z.read("word/document.xml").decode("utf-8", errors="replace")

    table_count = len(doc.tables)

    # Leftover markdown signals
    leftover_checks = {
        "raw ``` fence": "```" in full,
        "raw mermaid keyword as body": bool(re.search(r"(?m)^mermaid$", full)),
        "pipe-only table remnant lines": bool(re.search(r"(?m)^\s*\|[-:\s]+\|\s*$", full)),
        "unconverted ## heading marks": bool(re.search(r"(?m)^#{1,6}\s", full)),
    }

    # Figure captions
    fig_caps = re.findall(r"Figure\s+(\d+):", full)
    fig_nums = sorted(int(x) for x in fig_caps)

    # Approximate page count via app.xml or none
    page_count = "unknown"
    try:
        with zipfile.ZipFile(DOCX) as z:
            if "docProps/app.xml" in z.namelist():
                app = z.read("docProps/app.xml").decode("utf-8", errors="replace")
                m = re.search(r"<Pages>(\d+)</Pages>", app)
                if m:
                    page_count = int(m.group(1))
    except Exception:
        pass

    # Estimate pages from content length (rough)
    words = len(full.split())
    est_pages = max(1, words // 350 + image_count // 2)

    print("=== VERIFICATION ===")
    print(f"DOCX: {DOCX}")
    print(f"Size: {DOCX.stat().st_size} bytes")
    print(f"Paragraphs with text: {len(texts)}")
    print(f"Headings: {len(headings)}")
    print(f"Tables: {table_count}")
    print(f"Images (rels): {image_count}; media files: {len(media)}")
    print(f"Figure captions found: {len(fig_caps)} nums={fig_nums[:3]}...{fig_nums[-3:] if fig_nums else []}")
    print(f"Pages (app.xml): {page_count}; est~{est_pages}; words~{words}")
    print("\nSections:")
    for s, ok in section_status:
        print(f"  [{'OK' if ok else 'MISS'}] {s}")
    print("\nLeftover markdown:")
    for k, bad in leftover_checks.items():
        print(f"  [{'FAIL' if bad else 'OK'}] {k}: {bad}")

    # Cover / TOC markers
    cover_ok = "Complete Architecture and Flow Guide" in full and "Avinash258/contract-testing-business-demo" in full
    toc_ok = "Table of Contents" in full and "List of Diagrams" in full
    print(f"\nCover present: {cover_ok}")
    print(f"TOC + LoD present: {toc_ok}")

    # Count expected MD tables from earlier inventory = 60
    print(f"\nExpected figures: {len(META['figures'])}; actual images: {image_count}")
    print(f"Expected MD tables ~60; actual Word tables: {table_count}")

    checklist = [
        ("All 7 sections present", True, all(ok for _, ok in section_status), all(ok for _, ok in section_status)),
        ("Figure count = Mermaid count", len(META["figures"]), image_count, image_count == len(META["figures"])),
        ("Figure captions 1..N", len(META["figures"]), len(fig_caps), len(fig_caps) >= len(META["figures"])),
        ("Tables present (>=50)", ">=50", table_count, table_count >= 50),
        ("Cover page", True, cover_ok, cover_ok),
        ("TOC + List of Diagrams", True, toc_ok, toc_ok),
        ("No raw ```", False, leftover_checks["raw ``` fence"], not leftover_checks["raw ``` fence"]),
        ("No raw ## headings", False, leftover_checks["unconverted ## heading marks"], not leftover_checks["unconverted ## heading marks"]),
        ("No raw mermaid line", False, leftover_checks["raw mermaid keyword as body"], not leftover_checks["raw mermaid keyword as body"]),
    ]
    print("\nChecklist:")
    for name, exp, act, ok in checklist:
        print(f"  [{'PASS' if ok else 'FAIL'}] {name} | expected={exp} actual={act}")

    # Sample heading list
    print("\nFirst 25 headings:")
    for style, t in headings[:25]:
        print(f"  {style}: {t[:80]}")


if __name__ == "__main__":
    main()
